"""Word (.docx) engine.

* ``create_docx`` builds a document from the shared block model.
* ``fill_docx`` renders a Word template using docxtpl (Jinja2 tags inside Word).
* ``extract_docx_text`` reads the textual content back out.
"""

from __future__ import annotations

import io

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..markdown_parser import Block, Run

_MONO = "Consolas"


def _set_run_font(run, *, code: bool = False) -> None:
    if code:
        run.font.name = _MONO
        run.font.size = Pt(9)
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), _MONO)


def _shade(paragraph, fill: str = "F2F2F2") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _add_hyperlink(paragraph, url: str, text: str, bold: bool, italic: bool) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run_el = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    run_el.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run_el.append(t)
    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def _write_runs(paragraph, runs: list[Run]) -> None:
    for r in runs:
        if not r.text:
            continue
        if r.link:
            _add_hyperlink(paragraph, r.link, r.text, r.bold, r.italic)
            continue
        run = paragraph.add_run(r.text)
        run.bold = r.bold
        run.italic = r.italic
        _set_run_font(run, code=r.code)


def _add_block(doc: Document, b: Block) -> None:
    if b.type == "heading":
        level = max(1, min(b.level, 9))
        try:
            p = doc.add_paragraph(style=f"Heading {level}")
        except KeyError:
            p = doc.add_paragraph()
        _write_runs(
            p, [(Run(r.text, r.bold or True, r.italic, r.code, r.link)) for r in b.runs]
        )
    elif b.type == "paragraph":
        p = doc.add_paragraph()
        _write_runs(p, b.runs)
    elif b.type in ("bullet_list", "ordered_list"):
        base = "List Bullet" if b.type == "bullet_list" else "List Number"
        for item in b.items:
            style = base if item.depth == 0 else f"{base} {min(item.depth + 1, 3)}"
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph(style=base)
            _write_runs(p, item.runs)
    elif b.type == "code":
        lines = b.text.split("\n")
        p = doc.add_paragraph()
        _shade(p)
        for n, line in enumerate(lines):
            run = p.add_run(line)
            _set_run_font(run, code=True)
            if n < len(lines) - 1:
                run.add_break()
    elif b.type == "quote":
        try:
            p = doc.add_paragraph(style="Quote")
        except KeyError:
            p = doc.add_paragraph()
        _write_runs(p, b.runs)
        p.paragraph_format.left_indent = Pt(18)
    elif b.type == "table":
        _add_table(doc, b)
    elif b.type == "hr":
        p = doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "AAAAAA")
        borders.append(bottom)
        ppr.append(borders)


def _add_table(doc: Document, b: Block) -> None:
    cols = len(b.header) if b.header else (len(b.rows[0]) if b.rows else 0)
    if cols == 0:
        return
    rows = ([b.header] if b.header else []) + b.rows
    table = doc.add_table(rows=len(rows), cols=cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = row[j] if j < len(row) else ""
            if i == 0 and b.header:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
    doc.add_paragraph()


def create_docx(blocks: list[Block], meta: dict | None = None) -> bytes:
    meta = meta or {}
    doc = Document()
    for b in blocks:
        _add_block(doc, b)
    cp = doc.core_properties
    if meta.get("title"):
        cp.title = str(meta["title"])
    if meta.get("author"):
        cp.author = str(meta["author"])
    if meta.get("subject"):
        cp.subject = str(meta["subject"])
    if meta.get("keywords"):
        cp.keywords = str(meta["keywords"])
    if meta.get("comments"):
        cp.comments = str(meta["comments"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def fill_docx(template_path: str, context: dict, strict: bool = True) -> bytes:
    from docxtpl import DocxTemplate

    doc = DocxTemplate(template_path)
    if strict:
        missing = sorted(doc.get_undeclared_template_variables() - set(context.keys()))
        if missing:
            raise ValueError("Missing template variables: " + ", ".join(missing))
    doc.render(context)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def extract_docx_text(path: str, include_headers: bool = True) -> str:
    doc = Document(path)
    out: list[str] = []
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            if item.text.strip():
                out.append(item.text)
        elif isinstance(item, Table):
            for row in item.rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
    if include_headers:
        for section in doc.sections:
            for hf in (section.header, section.footer):
                for p in hf.paragraphs:
                    if p.text.strip():
                        out.append(p.text)
    return "\n".join(out)
