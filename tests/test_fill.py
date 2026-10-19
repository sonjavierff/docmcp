from __future__ import annotations

import pytest
from docx import Document

from docmcp.document import extract_text, fill_template


def test_fill_text_template(tmp_path, context):
    tpl = tmp_path / "tpl.md"
    tpl.write_text(
        "# {{ titulo }}\n\n"
        "{% for s in secciones %}## {{ s.titulo }}\n\n{{ s.cuerpo }}\n\n{% endfor %}"
        "{% if version %}Version: {{ version }}\n{% endif %}",
        encoding="utf-8",
    )
    out = tmp_path / "out.md"
    fill_template(str(tpl), str(out), context)
    rendered = out.read_text(encoding="utf-8")
    assert "# Manual Test" in rendered
    assert "## Intro" in rendered
    assert "Version: 1.0" in rendered


def test_text_template_strict_raises(tmp_path, context):
    tpl = tmp_path / "tpl.md"
    tpl.write_text("{{ noexiste }}", encoding="utf-8")
    with pytest.raises(Exception) as excinfo:
        fill_template(str(tpl), str(tmp_path / "out.md"), context)
    assert "noexiste" in str(excinfo.value)


def test_text_template_non_strict(tmp_path, context):
    tpl = tmp_path / "tpl.md"
    tpl.write_text("[{{ noexiste }}]", encoding="utf-8")
    out = tmp_path / "out.md"
    fill_template(str(tpl), str(out), context, strict=False)
    assert out.read_text(encoding="utf-8") == "[]"


def test_fill_docx_template(tmp_path, docx_template, context):
    out = tmp_path / "out.docx"
    result = fill_template(str(docx_template), str(out), context)
    assert result["format"] == "docx"

    d = Document(str(out))
    body = [p.text for p in d.paragraphs]
    assert any("MANUAL: Manual Test" in p for p in body)
    assert any("Autor: Javi" in p for p in body)
    assert any("Versión 1.0" in p for p in body)
    assert any("A (a1, a2)" in p for p in body)

    rows = [[c.text for c in row.cells] for row in d.tables[0].rows]
    assert rows[0] == ["Nombre", "Valor"]
    assert ["n1", "v1"] in rows
    assert ["n2", "v2"] in rows
    assert len(rows) == 3  # header + 2 data rows


def test_fill_odt_template(tmp_path, odt_template, context):
    out = tmp_path / "out.odt"
    result = fill_template(str(odt_template), str(out), context)
    assert result["format"] == "odt"
    text = extract_text(str(out))
    assert "MANUAL: Manual Test" in text
    assert "Manual Test split" in text  # placeholder split across spans
    assert "- Intro: Bienvenido" in text
    assert "Versión 1.0" in text
    assert "Categoria 1: A" in text
    assert "* a1" in text
    assert "* b1" in text
    assert "| n1 | v1" in text
    assert "| n2 | v2" in text


def test_fill_docx_rejects_other_output(tmp_path, docx_template, context):
    with pytest.raises(ValueError):
        fill_template(str(docx_template), str(tmp_path / "out.odt"), context)
