from __future__ import annotations

import pytest
from docx import Document
from odf.opendocument import OpenDocumentText
from odf.table import Table, TableCell, TableColumn, TableRow
from odf.text import P, Span


def build_docx_template(path) -> None:
    d = Document()
    d.add_heading("MANUAL: {{ titulo }}", level=1)
    d.add_paragraph("Autor: {{ autor }}")
    d.add_paragraph(
        "{% for s in secciones %}{{ s.titulo }}: {{ s.cuerpo }} | {% endfor %}"
    )
    d.add_paragraph("{% if version %}Versión {{ version }}{% endif %}")
    d.add_paragraph(
        "Categorias: {% for c in categorias %}{{ c.nombre }} ({{ c.items|join(', ') }}) {% endfor %}"
    )
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    h = t.rows[0].cells
    h[0].text = "Nombre"
    h[1].text = "Valor"
    marker = t.add_row().cells
    marker[0].text = "{%tr for r in filas %}"
    data = t.add_row().cells
    data[0].text = "{{ r.nombre }}"
    data[1].text = "{{ r.valor }}"
    end = t.add_row().cells
    end[0].text = "{%tr endfor %}"
    d.save(str(path))


def build_odt_template(path) -> None:
    doc = OpenDocumentText()

    def plain(text):
        p = P()
        p.addText(text)
        doc.text.addElement(p)

    def split(*chunks):
        p = P()
        for c in chunks:
            p.addElement(Span(text=c))
        doc.text.addElement(p)

    plain("MANUAL: {{ titulo }}")
    plain("Autor: {{ autor }}")
    split("{{ tit", "ulo }} split")  # placeholder split across spans
    plain("{% for s in secciones %}")
    plain("- {{ s.titulo }}: {{ s.cuerpo }}")
    plain("{% endfor %}")
    plain("{% if version %}")
    plain("Versión {{ version }}")
    plain("{% endif %}")
    plain("{% for c in categorias %}")
    plain("Categoria {{ loop.index }}: {{ c.nombre }}")
    plain("{% for i in c.items %}")
    plain("   * {{ i }}")
    plain("{% endfor %}")
    plain("{% endfor %}")

    tbl = Table(name="t1")
    tbl.addElement(TableColumn(numbercolumnsrepeated=3))
    tr = TableRow()
    for value in ("Nombre", "Valor", "Extra"):
        tc = TableCell(valuetype="string")
        p = P()
        p.addText(value)
        tc.addElement(p)
        tr.addElement(tc)
    tbl.addElement(tr)

    def row(cells):
        tr = TableRow()
        for value in cells:
            tc = TableCell(valuetype="string")
            p = P()
            p.addText(value)
            tc.addElement(p)
            tr.addElement(tc)
        tbl.addElement(tr)

    row(["{% for r in filas %}", "{{ r.nombre }}", "{{ r.valor }}"])
    row(["{% endfor %}", "", ""])
    doc.text.addElement(tbl)
    doc.save(str(path))


@pytest.fixture
def context() -> dict:
    return {
        "titulo": "Manual Test",
        "autor": "Javi",
        "version": "1.0",
        "secciones": [
            {"titulo": "Intro", "cuerpo": "Bienvenido"},
            {"titulo": "Uso", "cuerpo": "Pulsa aquí"},
        ],
        "categorias": [
            {"nombre": "A", "items": ["a1", "a2"]},
            {"nombre": "B", "items": ["b1"]},
        ],
        "filas": [
            {"nombre": "n1", "valor": "v1"},
            {"nombre": "n2", "valor": "v2"},
        ],
    }


@pytest.fixture
def docx_template(tmp_path):
    path = tmp_path / "tpl.docx"
    build_docx_template(path)
    return path


@pytest.fixture
def odt_template(tmp_path):
    path = tmp_path / "tpl.odt"
    build_odt_template(path)
    return path
